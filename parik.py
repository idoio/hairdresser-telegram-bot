import asyncio
import sys
import os
import re
import logging
from datetime import datetime, date, timedelta
from typing import List, Dict

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from aiogram import Bot, Dispatcher, F
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardMarkup, InlineKeyboardButton,
    Message, CallbackQuery, FSInputFile
)
from aiogram.filters import Command
from dotenv import load_dotenv

# ===== НАСТРОЙКИ СИСТЕМЫ =====
if sys.platform == 'win32':
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

load_dotenv()
BOT_TOKEN = os.getenv("BOT_TOKEN")
MASTER_CHAT_ID = int(os.getenv("MASTER_CHAT_ID", 0))
WORK_HOURS_START = 9
WORK_HOURS_END = 20
EXCEL_FILE = "appointments.xlsx"


# ===== УПРАВЛЕНИЕ EXCEL (С АВТОВОССТАНОВЛЕНИЕМ) =====
class HairdresserData:
    def __init__(self):
        if not os.path.exists(EXCEL_FILE):
            self._create_fresh_file()
            logger.info("✅ Создан appointments.xlsx")

    def _create_fresh_file(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Записи"
        ws.append(["Дата", "Время", "Имя клиента", "Телефон", "Статус", "user_id"])
        self._style_header(ws)
        wb.save(EXCEL_FILE)

    def _load_safe(self):
        try:
            return load_workbook(EXCEL_FILE)
        except Exception as e:
            logger.warning(f"⚠️ Файл Excel поврежден ({type(e).__name__}). Пересоздаю чистый...")
            if os.path.exists(EXCEL_FILE):
                try:
                    os.remove(EXCEL_FILE)
                except:
                    pass
            self._create_fresh_file()
            return load_workbook(EXCEL_FILE)

    def _style_header(self, ws):
        font = Font(bold=True, color="FFFFFF", size=11)
        fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
        border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
        for cell in ws[1]:
            cell.font, cell.fill, cell.border = font, fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def check_slot(self, d: str, t: str) -> bool:
        wb = self._load_safe()
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            if r[0] == d and r[1] == t and r[4] == "Активна":
                return False
        return True

    def save(self, d: str, t: str, name: str, phone: str, user_id: int) -> bool:
        wb = self._load_safe()
        ws = wb.active
        ws.append([d, t, name, phone, "Активна", user_id])
        wb.save(EXCEL_FILE)
        return True

    def cancel(self, date: str, time: str, phone: str) -> bool:
        wb = self._load_safe()
        for r in range(2, wb.active.max_row + 1):
            if (wb.active.cell(row=r, column=1).value == date and
                    wb.active.cell(row=r, column=2).value == time and
                    wb.active.cell(row=r, column=4).value == phone):
                wb.active.cell(row=r, column=5, value="Отменена")
                wb.save(EXCEL_FILE)
                return True
        return False

    def check_ownership(self, date: str, time: str, phone: str, telegram_id: int) -> bool:
        """Проверяет, принадлежит ли запись данному пользователю"""
        wb = self._load_safe()
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            # 0=Дата, 1=Время, 3=Телефон, 4=Статус, 5=user_id
            if r[0] == date and r[1] == time and r[3] == phone and r[5] == telegram_id and r[4] == "Активна":
                return True
        return False

    def get_user_data(self, telegram_id: int) -> Dict:
        wb = self._load_safe()
        last_found = None
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            if r[5] == telegram_id:
                last_found = {"name": r[2], "phone": r[3]}
        return last_found

    def get_active(self) -> List[Dict]:
        """Для мастера: все активные записи"""
        wb = self._load_safe()
        result = []
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            if r[4] == "Активна":
                result.append({"date": r[0], "time": r[1], "name": r[2], "phone": r[3]})
        result.sort(key=lambda x: (datetime.strptime(x["date"], "%d.%m.%Y"), x["time"]))
        return result

    def get_user_appointments(self, telegram_id: int) -> List[Dict]:
        """🔑 ИСПРАВЛЕНИЕ: Только записи текущего пользователя"""
        wb = self._load_safe()
        result = []
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            if r[5] == telegram_id and r[4] == "Активна":
                result.append({"date": r[0], "time": r[1], "name": r[2], "phone": r[3]})
        result.sort(key=lambda x: (datetime.strptime(x["date"], "%d.%m.%Y"), x["time"]))
        return result

    def get_by_date(self, target: str) -> List[Dict]:
        return [a for a in self.get_active() if a["date"] == target]


data_mgr = HairdresserData()


# ===== FSM & КЛАВИАТУРЫ =====
class Booking(StatesGroup):
    date = State()
    time = State()
    name = State()
    phone = State()
    confirm = State()


def kb_main():
    return ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text="📅 Записаться")],
        [KeyboardButton(text="📋 Мои записи"), KeyboardButton(text="❌ Отменить запись")],
        [KeyboardButton(text="👨‍💼 Панель мастера")]
    ], resize_keyboard=True)


def kb_dates():
    kb = []
    today = date.today()
    days_ru = {"Monday": "Пн", "Tuesday": "Вт", "Wednesday": "Ср", "Thursday": "Чт", "Friday": "Пт", "Saturday": "Сб",
               "Sunday": "Вс"}
    for i in range(7):
        dt = today + timedelta(days=i)
        date_str = f"{dt.day:02d}.{dt.month:02d}.{dt.year}"
        day_name = days_ru[dt.strftime('%A')]
        if i == 0:
            label = f"📅 Сегодня ({date_str})"
        elif i == 1:
            label = f"📅 Завтра ({date_str})"
        else:
            label = f"📅 {date_str} ({day_name})"
        kb.append([KeyboardButton(text=label)])
    kb.append([KeyboardButton(text="↩️ В главное меню")])
    return ReplyKeyboardMarkup(keyboard=kb, resize_keyboard=True)


def kb_times():
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text=f"{h:02d}:00")] for h in range(WORK_HOURS_START, WORK_HOURS_END)] + [
            [KeyboardButton(text="↩️ В главное меню")]], resize_keyboard=True)


def kb_back():
    return ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text="↩️ В главное меню")]], resize_keyboard=True)


def kb_phone():
    return ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text="📱 Отправить номер", request_contact=True)],
        [KeyboardButton(text="↩️ В главное меню")]
    ], resize_keyboard=True)


# ===== БОТ =====
dp = Dispatcher(storage=MemoryStorage())
bot = Bot(token=BOT_TOKEN)


@dp.message(Command("start"))
async def cmd_start(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer("👋 Добро пожаловать! Выберите действие:", reply_markup=kb_main())


@dp.message(F.text == "📅 Записаться")
async def start_booking(msg: Message, state: FSMContext):
    await state.clear()
    saved = data_mgr.get_user_data(msg.from_user.id)

    if saved:
        await state.update_data(name=saved['name'], phone=saved['phone'])
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Да, использовать", callback_data="use_saved")],
            [InlineKeyboardButton(text="✏️ Ввести новые данные", callback_data="enter_new")]
        ])
        await msg.answer(
            f"👋 Мы вас помним, <b>{saved['name']}</b>!\n\n"
            f"Использовать сохранённый номер: <code>{saved['phone']}</code>?",
            parse_mode="HTML", reply_markup=kb
        )
    else:
        await state.set_state(Booking.date)
        await msg.answer("📅 Выберите дату:", reply_markup=kb_dates())


@dp.callback_query(F.data.in_(["use_saved", "enter_new"]))
async def handle_saved_data(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    if callback.data == "use_saved":
        await state.set_state(Booking.date)
        await callback.message.answer("📅 Выберите дату:", reply_markup=kb_dates())
    else:
        await state.update_data(name=None, phone=None)
        await state.set_state(Booking.date)
        await callback.message.answer("📅 Выберите дату:", reply_markup=kb_dates())


@dp.message(Booking.date)
async def pick_date(msg: Message, state: FSMContext):
    if msg.text == "↩️ В главное меню":
        await state.clear()
        return await msg.answer("👋 Главное меню", reply_markup=kb_main())
    m = re.search(r'(\d{2}\.\d{2}\.\d{4})', msg.text)
    if not m: return await msg.answer("⚠️ Пожалуйста, выберите дату из списка")
    sel = datetime.strptime(m.group(1), "%d.%m.%Y").date()
    if sel < date.today(): return await msg.answer("❌ Нельзя выбрать прошедшую дату")
    await state.update_data(date=m.group(1))
    await state.set_state(Booking.time)
    await msg.answer("🕒 Выберите время:", reply_markup=kb_times())


@dp.message(Booking.time)
async def pick_time(msg: Message, state: FSMContext):
    if msg.text == "↩️ В главное меню":
        await state.clear()
        return await msg.answer("👋 Главное меню", reply_markup=kb_main())
    if not re.match(r'^\d{2}:00$', msg.text):
        return await msg.answer("⚠️ Выберите время из списка")

    data = await state.get_data()
    if not data_mgr.check_slot(data["date"], msg.text):
        return await msg.answer("❌ Это время уже занято. Выберите другое.")

    await state.update_data(time=msg.text)
    data = await state.get_data()  # Перечитываем после update

    if data.get("name") and data.get("phone"):
        await state.set_state(Booking.confirm)
        kb_confirm = ReplyKeyboardMarkup(keyboard=[
            [KeyboardButton(text="✅ Подтвердить")],
            [KeyboardButton(text="❌ Отменить")]
        ], resize_keyboard=True)
        await msg.answer(
            f"📋 Проверьте:\n📅 {data['date']} 🕒 {data['time']}\n👤 {data['name']}\n📞 {data['phone']}\nВсё верно?",
            reply_markup=kb_confirm
        )
    else:
        await state.set_state(Booking.name)
        await msg.answer("✍️ Введите ваше имя:", reply_markup=kb_back())


@dp.message(Booking.name)
async def pick_name(msg: Message, state: FSMContext):
    if msg.text == "↩️ В главное меню":
        await state.clear()
        return await msg.answer("👋 Главное меню", reply_markup=kb_main())
    if not msg.text or len(msg.text.strip()) < 2:
        return await msg.answer("⚠️ Имя слишком короткое. Введите заново:")
    await state.update_data(name=msg.text.strip())
    await state.set_state(Booking.phone)
    await msg.answer("📞 Введите номер или нажмите кнопку:", reply_markup=kb_phone())


@dp.message(Booking.phone)
async def pick_phone(msg: Message, state: FSMContext):
    if msg.text == "↩️ В главное меню":
        await state.clear()
        return await msg.answer("👋 Главное меню", reply_markup=kb_main())

    phone = ""
    if msg.contact:
        phone = msg.contact.phone_number
    elif msg.text:
        phone = msg.text.strip().replace(" ", "")
    else:
        return await msg.answer("⚠️ Пожалуйста, введите номер текстом или нажмите '📱 Отправить номер'.")

    if phone.startswith("8"):
        phone = "+7" + phone[1:]
    elif phone.startswith("7") and not phone.startswith("+"):
        phone = "+" + phone

    if not re.match(r'^\+7\d{10}$', phone):
        return await msg.answer("❌ Неверный формат. Используйте +79991234567")

    await state.update_data(phone=phone)
    await state.set_state(Booking.confirm)
    data = await state.get_data()

    kb_confirm = ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text="✅ Подтвердить")],
        [KeyboardButton(text="❌ Отменить")]
    ], resize_keyboard=True)
    await msg.answer(
        f"📋 Проверьте:\n📅 {data.get('date', '?')} 🕒 {data.get('time', '?')}\n👤 {data['name']}\n📞 {data['phone']}\nВсё верно?",
        reply_markup=kb_confirm)


@dp.message(Booking.confirm)
async def confirm_booking(msg: Message, state: FSMContext):
    if msg.text == "❌ Отменить":
        await state.clear()
        return await msg.answer("✅ Запись отменена.", reply_markup=kb_main())
    if msg.text != "✅ Подтвердить":
        return await msg.answer("⚠️ Выберите вариант из меню.")

    data = await state.get_data()
    success = data_mgr.save(data["date"], data["time"], data["name"], data["phone"], msg.from_user.id)
    if success:
        text = f"🎉 Записано!\n📅 {data['date']} в {data['time']}"
        await msg.answer(text, reply_markup=kb_main())
        try:
            await bot.send_message(MASTER_CHAT_ID,
                                   f"🔔 НОВАЯ ЗАПИСЬ\n📅 {data['date']} {data['time']}\n👤 {data['name']}\n📞 {data['phone']}")
            await bot.send_document(MASTER_CHAT_ID, FSInputFile(EXCEL_FILE), caption="📁 База обновлена")
        except Exception as e:
            logger.error(f"❌ Ошибка отправки мастеру: {e}")
    else:
        await msg.answer("❌ Ошибка сохранения.")
    await state.clear()


# ===== ПАНЕЛЬ МАСТЕРА & ОТМЕНА =====
@dp.message(F.text == "👨‍💼 Панель мастера")
async def master_panel(msg: Message, state: FSMContext):
    if msg.from_user.id != MASTER_CHAT_ID:
        return await msg.answer("🔒 Доступ только для мастера.")
    await state.clear()
    today = date.today().strftime("%d.%m.%Y")
    tmr = (date.today() + timedelta(days=1)).strftime("%d.%m.%Y")
    today_apps = data_mgr.get_by_date(today)
    tmr_apps = data_mgr.get_by_date(tmr)

    def fmt(title, apps):
        if not apps: return f"\n📅 <b>{title}:</b> записей нет"
        res = f"\n📅 <b>{title}:</b>\n"
        for a in apps: res += f"  🕒 {a['time']} | {a['name']} ({a['phone']})\n"
        return res

    text = f"📊 <b>Расписание</b>{fmt('Сегодня', today_apps)}{fmt('Завтра', tmr_apps)}"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 Скачать Excel", callback_data="dl_excel")],
        [InlineKeyboardButton(text="📋 Все записи", callback_data="show_all")]
    ])
    await msg.answer(text, parse_mode="HTML", reply_markup=kb)


@dp.callback_query(F.data.in_(["dl_excel", "show_all"]))
async def master_cb(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    if callback.from_user.id != MASTER_CHAT_ID: return await callback.answer("🔒 Нет доступа")
    if callback.data == "dl_excel":
        await bot.send_document(callback.message.chat.id, FSInputFile(EXCEL_FILE), caption="📁 Полная база")
    elif callback.data == "show_all":
        apps = data_mgr.get_active()
        if not apps: return await callback.answer("📝 Записей нет", show_alert=True)
        txt = "📋 <b>Все активные записи:</b>\n" + "\n".join(
            f"🕒 {a['time']} | {a['date']} | {a['name']} | {a['phone']}" for a in apps
        )
        for i in range(0, len(txt), 4000): await callback.message.answer(txt[i:i + 4000], parse_mode="HTML")


# 🔑 ИСПРАВЛЕНИЕ: Безопасная отмена только своих записей
@dp.callback_query(F.data.startswith("cancel_"))
async def cancel_cb(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    parts = callback.data.split("_", 3)
    if len(parts) != 4: return await callback.answer("❌ Ошибка формата", show_alert=True)
    _, date, time, phone_safe = parts
    phone = phone_safe.replace("p", "+")

    # Проверка владельца перед отменой
    if not data_mgr.check_ownership(date, time, phone, callback.from_user.id):
        return await callback.answer("🔒 Это не ваша запись!", show_alert=True)

    if data_mgr.cancel(date, time, phone):
        await callback.message.edit_text(f"🗑 Ваша запись на {date} {time} успешно отменена.")
        try:
            await bot.send_message(MASTER_CHAT_ID,
                                   f"🗑 КЛИЕНТ ОТМЕНИЛ ЗАПИСЬ\n📅 {date} {time}\n👤 {callback.from_user.username or 'Клиент'}")
        except:
            pass
    else:
        await callback.answer("❌ Ошибка или уже отменена", show_alert=True)


# 🔑 ИСПРАВЛЕНИЕ: Показ только записей текущего пользователя
@dp.message(F.text == "📋 Мои записи")
async def my_apps(msg: Message, state: FSMContext):
    await state.clear()
    apps = data_mgr.get_user_appointments(msg.from_user.id)

    if not apps:
        return await msg.answer("📝 У вас пока нет записей.", reply_markup=kb_main())

    text = "📋 Ваши записи:\n" + "\n".join(f"🕒 {a['time']} | {a['date']} | {a['name']}" for a in apps)
    kb = InlineKeyboardMarkup(inline_keyboard=[])
    for a in apps:
        phone_safe = a['phone'].replace('+', 'p')
        kb.inline_keyboard.append([InlineKeyboardButton(text=f"❌ Отменить {a['date']} {a['time']}",
                                                        callback_data=f"cancel_{a['date']}_{a['time']}_{phone_safe}")])
    await msg.answer(text, reply_markup=kb)


@dp.message(F.text == "❌ Отменить запись")
async def cancel_menu(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer("Выберите запись в разделе 📋 Мои записи, чтобы отменить её.", reply_markup=kb_main())


async def main():
    logger.info("🚀 Запуск бота для парикмахера...")
    await dp.start_polling(bot, skip_updates=True)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        logger.info("🛑 Остановлен")